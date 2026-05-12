#!/usr/bin/env python3
"""
导出 Word 投标响应文档
将响应内容导出为规范的 Word 文档。
"""

import json
import sys
from pathlib import Path
from typing import Dict, Any, List

try:
    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("请安装 python-docx: pip install python-docx")
    sys.exit(1)


# 文档样式配置
STYLES = {
    'title': {
        'font_size': Pt(22),
        'bold': True,
        'alignment': WD_ALIGN_PARAGRAPH.CENTER
    },
    'heading1': {
        'font_size': Pt(16),
        'bold': True,
        'space_before': Pt(12),
        'space_after': Pt(6)
    },
    'heading2': {
        'font_size': Pt(14),
        'bold': True,
        'space_before': Pt(6),
        'space_after': Pt(3)
    },
    'body': {
        'font_size': Pt(12),
        'space_after': Pt(3)
    },
    'table_header': {
        'font_size': Pt(11),
        'bold': True,
        'bg_color': RGBColor(220, 220, 220)
    },
    'table_body': {
        'font_size': Pt(10)
    }
}


def load_response_content(file_path: str) -> Dict[str, Any]:
    """加载响应内容"""
    with open(file_path, 'r', encoding='utf-8') as f:
        return json.load(f)


def add_title(doc: Document, text: str) -> None:
    """添加文档标题"""
    title = doc.add_heading(text, 0)
    title.alignment = STYLES['title']['alignment']


def add_heading1(doc: Document, text: str) -> None:
    """添加一级标题"""
    heading = doc.add_heading(text, 1)
    return heading


def add_heading2(doc: Document, text: str) -> None:
    """添加二级标题"""
    heading = doc.add_heading(text, 2)
    return heading


def add_paragraph(doc: Document, text: str, style: str = 'body') -> None:
    """添加段落"""
    para = doc.add_paragraph(text)
    para.style.font.size = STYLES[style]['font_size']
    if 'space_after' in STYLES[style]:
        para.paragraph_format.space_after = STYLES[style]['space_after']
    return para


def add_section_table(doc: Document, sections: List[Dict[str, Any]], title: str) -> None:
    """添加章节表格"""
    add_heading2(doc, title)

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'

    headers = ["序号", "招标要求", "响应内容"]
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        cell.paragraphs[0].runs[0].bold = True

    for idx, section in enumerate(sections):
        row = table.add_row()
        row.cells[0].text = str(idx + 1)
        row.cells[1].text = section.get('requirement', '')
        row.cells[2].text = section.get('response', '')


def add_scoring_table(doc: Document, scoring: List[Dict[str, Any]]) -> None:
    """添加打分表"""
    add_heading2(doc, '评分标准响应')

    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'

    headers = ["评分项", "招标要求", "响应内容", "证明材料"]
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        cell.paragraphs[0].runs[0].bold = True

    for item in scoring:
        row = table.add_row()
        row.cells[0].text = item.get('score', '')
        row.cells[1].text = item.get('requirement', '')
        row.cells[2].text = item.get('response', '')
        row.cells[3].text = item.get('proof', '')


def add_deviation_table(doc: Document, deviation_data: Dict[str, Any]) -> None:
    """添加偏离表"""
    add_heading1(doc, '技术/商务偏离表')

    # 技术偏离表
    add_heading2(doc, '一、技术偏离表')
    tech_table = doc.add_table(rows=1, cols=4)
    tech_table.style = 'Table Grid'

    headers = ["序号", "招标要求", "响应说明", "偏离情况"]
    for i, header in enumerate(headers):
        tech_table.rows[0].cells[i].text = header

    for idx, item in enumerate(deviation_data.get('technical', [])):
        row = tech_table.add_row()
        row.cells[0].text = str(idx + 1)
        row.cells[1].text = item.get('requirement', '')
        row.cells[2].text = item.get('response', '')
        row.cells[3].text = item.get('deviation', '')

    # 商务偏离表
    add_heading2(doc, '二、商务偏离表')
    biz_table = doc.add_table(rows=1, cols=4)
    biz_table.style = 'Table Grid'

    for i, header in enumerate(headers):
        biz_table.rows[0].cells[i].text = header

    for idx, item in enumerate(deviation_data.get('business', [])):
        row = biz_table.add_row()
        row.cells[0].text = str(idx + 1)
        row.cells[1].text = item.get('requirement', '')
        row.cells[2].text = item.get('response', '')
        row.cells[3].text = item.get('deviation', '')


def export_word(response_file: str, output_file: str, deviation_file: str = None) -> None:
    """导出 Word 文档"""
    content = load_response_content(response_file)

    doc = Document()

    # 标题
    add_title(doc, content.get('project_name', '投标响应文件'))

    # 项目信息
    add_heading1(doc, '一、项目概述')
    add_paragraph(doc, f"项目名称：{content.get('project_name', 'N/A')}")

    # 技术响应
    add_heading1(doc, '二、技术方案')

    tech_sections = [s for s in content.get('sections', []) if s.get('title') == '技术响应']
    if tech_sections:
        add_section_table(doc, tech_sections, '技术要求响应')

    # 商务响应
    add_heading1(doc, '三、商务方案')

    biz_sections = [s for s in content.get('sections', []) if s.get('title') == '商务响应']
    if biz_sections:
        add_section_table(doc, biz_sections, '商务要求响应')

    # 售后响应
    add_heading1(doc, '四、售后服务')

    after_sections = [s for s in content.get('sections', []) if s.get('title') == '售后响应']
    if after_sections:
        add_section_table(doc, after_sections, '售后要求响应')

    # 打分表
    if content.get('scoring'):
        add_heading1(doc, '五、评分标准响应')
        add_scoring_table(doc, content['scoring'])

    # 偏离表
    if deviation_file:
        with open(deviation_file, 'r', encoding='utf-8') as f:
            deviation_data = json.load(f)['data']
        add_deviation_table(doc, deviation_data)

    doc.save(output_file)
    print(f"Word 文档已生成: {output_file}")


def main():
    if len(sys.argv) < 3:
        print("用法: python export_word.py <响应内容.json> <输出文件.docx> [--deviation <偏离表.json>]")
        sys.exit(1)

    response_file = sys.argv[1]
    output_file = sys.argv[2]

    deviation_file = None
    for i, arg in enumerate(sys.argv):
        if arg == '--deviation' and i + 1 < len(sys.argv):
            deviation_file = sys.argv[i + 1]

    export_word(response_file, output_file, deviation_file)


if __name__ == "__main__":
    main()