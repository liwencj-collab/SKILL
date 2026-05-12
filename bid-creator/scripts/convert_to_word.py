#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# 读取Markdown文件
md_file = '/Users/qujinyu/Downloads/qiuzhi-skill-creator/bid-creator/技术投标文件-按格式.md'
docx_file = '/Users/qujinyu/Downloads/qiuzhi-skill-creator/bid-creator/技术投标文件-按格式.docx'

with open(md_file, 'r', encoding='utf-8') as f:
    content = f.read()

# 创建Word文档
doc = Document()

# 设置中文字体
def set_chinese_font(run, font_name='宋体', font_size=12):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

# 解析Markdown并添加到文档
lines = content.split('\n')

i = 0
while i < len(lines):
    line = lines[i].strip()

    # 跳过空行
    if not line:
        i += 1
        continue

    # 标题处理
    if line.startswith('# '):
        p = doc.add_heading(line[2:], level=1)
        for run in p.runs:
            set_chinese_font(run, '黑体', 22)
    elif line.startswith('## '):
        p = doc.add_heading(line[3:], level=2)
        for run in p.runs:
            set_chinese_font(run, '黑体', 18)
    elif line.startswith('### '):
        p = doc.add_heading(line[4:], level=3)
        for run in p.runs:
            set_chinese_font(run, '黑体', 15)
    elif line.startswith('#### '):
        p = doc.add_heading(line[5:], level=4)
        for run in p.runs:
            set_chinese_font(run, '黑体', 14)
    elif line.startswith('##### '):
        p = doc.add_heading(line[6:], level=5)
        for run in p.runs:
            set_chinese_font(run, '黑体', 13)
    elif line.startswith('###### '):
        p = doc.add_heading(line[7:], level=6)
        for run in p.runs:
            set_chinese_font(run, '宋体', 12)
    else:
        # 处理表格 - 简化处理
        if line.startswith('|'):
            # 收集表格行
            table_lines = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1

            if len(table_lines) >= 2:
                # 获取表头
                header_parts = [cell.strip() for cell in table_lines[0].split('|')[1:-1]]
                if not header_parts:
                    i += 1
                    continue

                # 创建表格
                table = doc.add_table(rows=min(len(table_lines)-2, 20), cols=len(header_parts))
                table.style = 'Table Grid'

                # 设置表头
                for j, header in enumerate(header_parts):
                    if j < len(header_parts):
                        cell = table.rows[0].cells[j]
                        cell.text = header
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                set_chinese_font(run, '宋体', 9)
                                run.font.bold = True

                # 设置数据行（最多20行）
                for row_idx in range(2, min(len(table_lines), 22)):
                    row_parts = [cell.strip() for cell in table_lines[row_idx].split('|')[1:-1]]
                    for col_idx, part in enumerate(row_parts):
                        if col_idx < len(header_parts) and row_idx-1 < len(table.rows):
                            cell = table.rows[row_idx-1].cells[col_idx]
                            cell.text = part
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    set_chinese_font(run, '宋体', 9)
                continue
        else:
            # 处理代码块
            if line.startswith('```'):
                code_lines = []
                i += 1
                while i < len(lines) and not lines[i].strip().startswith('```'):
                    code_lines.append(lines[i])
                    i += 1

                p = doc.add_paragraph()
                for code_line in code_lines:
                    run = p.add_run(code_line + '\n')
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9)
                continue
            else:
                # 普通段落 - 处理加粗
                p = doc.add_paragraph()
                parts = re.split(r'(\*\*[^*]+\*\*)', line)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        run = p.add_run(part[2:-2])
                        run.font.bold = True
                        set_chinese_font(run, '宋体', 12)
                    else:
                        run = p.add_run(part)
                        set_chinese_font(run, '宋体', 12)

    i += 1

# 保存文档
doc.save(docx_file)
print(f"Word文档已生成: {docx_file}")
